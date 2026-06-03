#!/usr/bin/env python3
"""files-request: ベース IRL Excel をコピーし、PJ 用にファイル名と PJ 名称を書き換える。

やることは極めてシンプル：
  1. ./3_KnowledgeSource/questions-for-<business-model>.xlsx を
  2. ./4_Task/<yyyymmdd>-RequestFilesv<version>/<PJ名>_RequestFiles_v<version>.xlsx に
     コピーし、
  3. コピー先 Excel の全セルからプレースホルダ文字列を PJ 名に置換する。

中間フォーマット化・項目の再整形・管理列付与などは一切行わない（忠実コピー + 置換のみ）。

Usage:
  python make_request_files.py \\
    --business-model stores \\
    --pj-name "対象会社名" \\
    --placeholder "[PJ名]" \\
    --version 1

cwd は PJ ルートフォルダである前提（I/O はすべて cwd 相対）。
"""
from __future__ import annotations

import argparse
import datetime
import shutil
import sys
from pathlib import Path

# ビジネスモデル 5 分類の機械的 enum。
# 唯一の正本は references/business_model_taxonomy.md（ここはそれに従う写し）。
BUSINESS_MODELS = {
    "stores": "店舗・拠点型",
    "assets": "資産型",
    "labors": "労働型",
    "products": "製品型",
    "flows": "仲介・流通型",
}

KNOWLEDGE_DIR = "3_KnowledgeSource"
TASK_DIR = "4_Task"


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument(
        "--business-model",
        required=True,
        help="ビジネスモデル分類: " + " / ".join(BUSINESS_MODELS),
    )
    p.add_argument(
        "--pj-name",
        required=True,
        help="プロジェクト名（ファイル名 + Excel 内置換に使用）",
    )
    p.add_argument(
        "--placeholder",
        default=None,
        help="Excel 内で PJ 名に置換するプレースホルダ文字列（例: [PJ名]）。"
        "未指定なら置換せず純粋コピー。",
    )
    p.add_argument(
        "--version",
        type=int,
        default=1,
        help="リクエストファイルのバージョン番号（既定 1）",
    )
    p.add_argument(
        "--date",
        default=None,
        help="出力フォルダの日付プレフィックス YYYYMMDD（既定は本日）",
    )
    p.add_argument(
        "--root",
        default=".",
        help="PJ ルートフォルダ（既定: cwd）",
    )
    return p.parse_args()


def fail(msg: str) -> "NoReturn":  # type: ignore[name-defined]
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(1)


def replace_placeholder_in_workbook(path: Path, placeholder: str, pj_name: str) -> int:
    """全シートの全セルからプレースホルダ文字列を PJ 名に置換し、置換セル数を返す。"""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    count = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and placeholder in cell.value:
                    cell.value = cell.value.replace(placeholder, pj_name)
                    count += 1
    if count:
        wb.save(path)
    return count


def main() -> int:
    args = parse_args()

    # 1) ビジネスモデルの検証（未定義は捏造せず停止）
    bm = args.business_model
    if bm not in BUSINESS_MODELS:
        fail(
            f"ビジネスモデル '{bm}' は定義されていません。"
            f"有効な値: {', '.join(BUSINESS_MODELS)}"
        )

    root = Path(args.root).resolve()

    # 2) ベース IRL の存在確認（無ければ停止）
    base = root / KNOWLEDGE_DIR / f"questions-for-{bm}.xlsx"
    if not base.exists():
        fail(
            f"ベース IRL が見つかりません: {base}\n"
            f"  '{bm}'（{BUSINESS_MODELS[bm]}）用の questions-for-{bm}.xlsx を "
            f"{KNOWLEDGE_DIR}/ に配置してください。"
        )

    # 3) 出力先の決定
    if args.date:
        date_str = args.date
    else:
        date_str = datetime.date.today().strftime("%Y%m%d")
    out_dir = root / TASK_DIR / f"{date_str}-RequestFilesv{args.version}"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"{args.pj_name}_RequestFiles_v{args.version}.xlsx"

    # 4) コピー（書式・シート構成をそのまま保持）
    shutil.copy2(base, out_path)

    # 5) PJ 名の置換
    replaced = 0
    if args.placeholder:
        replaced = replace_placeholder_in_workbook(
            out_path, args.placeholder, args.pj_name
        )

    # 6) 報告
    print(f"OK: {out_path}")
    print(f"  base        : {base}")
    print(f"  business    : {bm}（{BUSINESS_MODELS[bm]}）")
    print(f"  pj_name     : {args.pj_name}")
    if args.placeholder:
        print(f"  placeholder : '{args.placeholder}' → '{args.pj_name}'（{replaced} セル置換）")
    else:
        print("  placeholder : なし（純粋コピー）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
